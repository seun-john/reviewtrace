"""Load a DOCX into the normalised :class:`~reviewtrace.models.document.Document`.

python-docx is used to open and validate the package and to resolve style names. Body
structure, comment ranges and tracked revisions come from :mod:`.ooxml`, which walks the
same element tree once so that every locator is consistent.
"""

from __future__ import annotations

import re
from pathlib import Path

from reviewtrace.extractors import ooxml
from reviewtrace.models.document import (
    Document,
    Footnote,
    Paragraph,
    Reference,
    Section,
    Table,
    TrackedChange,
    TrackedKind,
)
from reviewtrace.utils.errors import DocumentReadError, UnsupportedFileTypeError
from reviewtrace.utils.text import normalize

_HEADING_STYLE = re.compile(r"^heading\s*(\d+)$", re.I)
_NUMBERED = re.compile(r"^\s*(\d+(?:\.\d+)*)[.)]?\s+(\S.*)$")
_CHAPTER = re.compile(r"^\s*chapter\s+(\d+)\s*[:.\-–—]?\s*(.*)$", re.I)
_CAPTION = re.compile(
    r"^\s*(table|tab\.|figure|fig\.)\s*([A-Za-z]?\d+(?:[.\-]\d+)*)\b[.:)]?\s*(.*)$", re.I
)
_REF_HEADING = re.compile(
    r"^(references?( list)?|bibliography|works cited|literature cited|sources|reference)$", re.I
)
_YEAR_PAREN = re.compile(r"\(\s*((?:19|20)\d{2})[a-z]?\s*\)")
_YEAR_ANY = re.compile(r"\b((?:19|20)\d{2})[a-z]?\b")
_LEADING_LIST = re.compile(r"^\s*(?:\[\d+\]|\d+[.)]|[-•*])\s+")

SUPPORTED_SUFFIXES = {".docx"}


def check_docx_path(path: Path) -> None:
    if not path.exists():
        raise DocumentReadError(f"file not found: {path}")
    if not path.is_file():
        raise DocumentReadError(f"not a file: {path}")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        hint = " Convert legacy .doc files to .docx first." if path.suffix.lower() == ".doc" else ""
        raise UnsupportedFileTypeError(
            f"unsupported file type '{path.suffix or '(none)'}' for {path.name}; expected .docx.{hint}"
        )


def caption_info(text: str) -> tuple[str, str, str] | None:
    """Return ``(kind, number, title)`` for ``Table 4.2 ...`` / ``Figure 3 ...`` lines."""
    m = _CAPTION.match(text.strip())
    if not m:
        return None
    kind = "table" if m.group(1).lower().startswith("tab") else "figure"
    return kind, m.group(2), m.group(3).strip()


def parse_heading_text(text: str) -> tuple[str | None, str]:
    m = _CHAPTER.match(text)
    if m:
        return m.group(1), m.group(2).strip() or text.strip()
    m = _NUMBERED.match(text)
    if m:
        return m.group(1), m.group(2).strip()
    return None, text.strip()


def parse_year(raw: str) -> int | None:
    m = _YEAR_PAREN.search(raw) or _YEAR_ANY.search(raw)
    return int(m.group(1)) if m else None


def _all_bold(p_el: ooxml.Element) -> bool:
    runs = [
        r for r in p_el.iter(ooxml.q(ooxml.W, "r")) if "".join(str(t) for t in r.itertext()).strip()
    ]
    if not runs:
        return False
    for r in runs:
        rpr = r.find(ooxml.q(ooxml.W, "rPr"))
        if rpr is None or rpr.find(ooxml.q(ooxml.W, "b")) is None:
            return False
    return True


def load_document(
    path: str | Path, *, comments: bool = True, numbered_heading_fallback: bool = True
) -> Document:
    """Read ``path`` without modifying it.

    ``comments=False`` skips reading and anchoring Word comments (not needed for a document that
    is only used as a comparison target).
    ``numbered_heading_fallback`` treats short numbered lines as headings when a document
    has no styled headings; turn it off for reviewer reports, where numbered lines are items.
    """
    p = Path(path)
    check_docx_path(p)
    try:
        import docx
        from docx.enum.style import WD_STYLE_TYPE

        pkg_doc = docx.Document(str(p))
    except UnsupportedFileTypeError:
        raise
    except Exception as exc:  # python-docx raises several unrelated types for bad packages
        raise DocumentReadError(
            f"{p.name}: not a readable DOCX ({type(exc).__name__}: {exc})"
        ) from exc

    body = pkg_doc.element.body
    story = ooxml.walk_story(body)

    style_names: dict[str, str] = {}

    def style_name(style_id: str | None) -> str:
        if not style_id:
            return ""
        if style_id not in style_names:
            try:
                style = pkg_doc.styles.get_by_id(style_id, WD_STYLE_TYPE.PARAGRAPH)
                style_names[style_id] = style.name if style is not None and style.name else style_id
            except Exception:
                style_names[style_id] = style_id
        return style_names[style_id]

    body_raw = [rp for rp in story.paragraphs if rp.table is None and rp.number]

    def heading_level(rp: ooxml.RawParagraph, sname: str) -> int:
        m = _HEADING_STYLE.match(sname.strip()) or _HEADING_STYLE.match((rp.style_id or "").strip())
        if m:
            return int(m.group(1))
        if rp.outline_level is not None and rp.outline_level < 9:
            return rp.outline_level + 1
        return 0

    levels = {id(rp): heading_level(rp, style_name(rp.style_id)) for rp in body_raw}
    if numbered_heading_fallback and not any(levels.values()):
        # No structured headings at all: fall back to short numbered lines ("2.3 Sampling").
        for rp in body_raw:
            text = rp.text.strip()
            m = _NUMBERED.match(text)
            if (
                m
                and len(text) <= 120
                and not text.endswith((".", ",", ";", ":"))
                and (_all_bold(rp.element) or len(text.split()) <= 10)
            ):
                levels[id(rp)] = m.group(1).count(".") + 1

    doc = Document(path=str(p))
    doc.title = (pkg_doc.core_properties.title or "").strip()

    sections: list[Section] = [Section(id="S0", heading="", title="", level=0)]
    stack: list[Section] = []
    current = sections[0]
    paragraphs: list[Paragraph] = []
    for rp in body_raw:
        sname = style_name(rp.style_id)
        level = levels[id(rp)]
        text = rp.text.strip()
        if level:
            number, title = parse_heading_text(text)
            while stack and stack[-1].level >= level:
                stack.pop()
            current = Section(
                id=f"S{len(sections)}",
                heading=text,
                number=number,
                title=title,
                level=level,
                parent_id=stack[-1].id if stack else None,
                heading_paragraph=rp.number,
            )
            sections.append(current)
            stack.append(current)
        elif not doc.title and sname.lower() == "title":
            doc.title = text
        para = Paragraph(
            number=rp.number,
            text=text,
            style=sname,
            is_heading=bool(level),
            heading_level=level,
            is_list_item=rp.is_list or sname.lower().startswith("list"),
            has_image=rp.has_image,
            section_id=current.id,
        )
        current.paragraph_numbers.append(rp.number)
        paragraphs.append(para)
    doc.paragraphs = paragraphs
    doc.sections = sections
    doc.reindex()

    by_number = {p_.number: p_ for p_ in paragraphs}
    for raw_table in story.tables:
        after = 0
        if raw_table.after is not None:
            after = raw_table.after.number or raw_table.after.prev_number
        sec_id = by_number[after].section_id if after in by_number else "S0"
        rows = [[c.strip() for c in row] for row in raw_table.rows]
        table = Table(index=raw_table.index, rows=rows, section_id=sec_id, after_paragraph=after)
        _attach_caption(table, by_number)
        doc.tables.append(table)
        owner = doc.section(sec_id)
        if owner is not None:
            owner.table_indices.append(table.index)

    doc.tracked_changes = _tracked_changes(story, doc)
    doc.references = _references(doc)

    with ooxml.OoxmlPackage(p) as pkg:
        note_stories = _note_stories(pkg)
        _load_notes(doc, note_stories)
        if comments:
            from reviewtrace.extractors.comments import build_comments

            raw_comments = ooxml.parse_comments(pkg)
            doc.comments = build_comments(raw_comments, story, note_stories, doc)
    return doc


def _attach_caption(table: Table, by_number: dict[int, Paragraph]) -> None:
    """A caption is the paragraph just before the table, or failing that just after it."""
    for offset in (0, 1):
        para = by_number.get(table.after_paragraph + offset)
        if para is None:
            continue
        info = caption_info(para.text)
        if info and info[0] == "table":
            table.caption = para.text
            table.label = f"Table {info[1]}"
            return


def _tracked_changes(story: ooxml.Story, doc: Document) -> list[TrackedChange]:
    out: list[TrackedChange] = []
    for t in story.tracked:
        rp = t.paragraph
        number = rp.number or rp.prev_number
        table_index = cell = None
        if rp.table is not None:
            table_index, cell = rp.table[0], (rp.table[1], rp.table[2])
            table = doc.tables[table_index] if table_index < len(doc.tables) else None
            section_id = table.section_id if table else None
            number = table.after_paragraph if table else 0
        else:
            sec = doc.section_of(number) if number else doc.section("S0")
            section_id = sec.id if sec else None
        out.append(
            TrackedChange(
                id=t.id,
                kind=TrackedKind(t.kind),
                author=t.author,
                date=t.date,
                text=t.text.strip(),
                moved=t.moved,
                paragraph=number or None,
                section_id=section_id,
                table_index=table_index,
                cell=cell,
            )
        )
    return out


def _note_stories(pkg: ooxml.OoxmlPackage) -> list[tuple[str, ooxml.Story]]:
    return ooxml.parse_notes(pkg, "word/footnotes.xml", "footnote") + ooxml.parse_notes(
        pkg, "word/endnotes.xml", "endnote"
    )


def _load_notes(doc: Document, stories: list[tuple[str, ooxml.Story]]) -> None:
    for nid, st in stories:
        text = " ".join(p.text.strip() for p in st.paragraphs if p.text.strip())
        kind = "footnote" if st.name.startswith("footnote") else "endnote"
        if text:
            doc.footnotes.append(Footnote(id=nid, kind=kind, text=text))
        for t in st.tracked:
            doc.tracked_changes.append(
                TrackedChange(
                    id=t.id,
                    kind=TrackedKind(t.kind),
                    author=t.author,
                    date=t.date,
                    text=t.text.strip(),
                    moved=t.moved,
                    story=st.name,
                )
            )


def _references(doc: Document) -> list[Reference]:
    ref_sections = [
        s for s in doc.sections if s.title and _REF_HEADING.match(normalize(s.title).strip(" :"))
    ]
    out: list[Reference] = []
    seen_sections: set[str] = set()
    for sec in ref_sections:
        for sid in doc.subtree_ids(sec.id):
            if sid in seen_sections:
                continue
            seen_sections.add(sid)
            for para in doc.body_paragraphs(sid):
                raw = _LEADING_LIST.sub("", para.text).strip()
                if len(raw.split()) < 3:
                    continue
                out.append(
                    Reference(index=len(out), raw=raw, year=parse_year(raw), paragraph=para.number)
                )
    return out
